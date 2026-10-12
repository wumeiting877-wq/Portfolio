from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
FORBIDDEN_HEADERS = {
    "password",
    "password_hash",
    "old_password",
    "new_password",
    "initial_password",
    "api_key",
    "secret",
    "token",
    "credential",
    "access_key",
    "private_key",
    "ssh_private_key",
}

IMPORT_TYPES = (
    "employees",
    "projects",
    "project_members",
    "resource_types",
    "resources",
    "resource_allocations",
    "expenses",
)

TEMPLATE_COLUMNS: dict[str, list[tuple[str, str, str, str, str]]] = {
    # header, required, example, note, field
    "employees": [
        ("employee_no", "是", "E1001", "工号，唯一", "employee_no"),
        ("name", "是", "张三", "姓名", "name"),
        ("department", "否", "研发部", "部门", "department"),
        ("position", "否", "工程师", "职位", "position"),
        ("email", "否", "zhangsan@example.com", "邮箱格式", "email"),
        ("phone", "否", "13800000000", "电话", "phone"),
        ("employment_status", "否", "ACTIVE", "ACTIVE / INACTIVE / LEFT，默认 ACTIVE", "employment_status"),
        ("remark", "否", "", "备注", "remark"),
    ],
    "projects": [
        ("project_code", "是", "P-001", "项目编码，唯一", "project_code"),
        ("name", "是", "核心平台", "项目名称", "name"),
        ("manager_employee_no", "是", "E1001", "负责人必须存在且在职", "manager_employee_no"),
        ("status", "否", "ACTIVE", "PLANNING / ACTIVE / PAUSED / COMPLETED", "status"),
        ("description", "否", "", "说明", "description"),
        ("start_date", "否", "2026-01-01", "YYYY-MM-DD", "start_date"),
        ("end_date", "否", "2026-12-31", "YYYY-MM-DD", "end_date"),
    ],
    "project_members": [
        ("project_code", "是", "P-001", "已存在项目", "project_code"),
        ("employee_no", "是", "E1002", "已存在且在职人员", "employee_no"),
        ("member_role", "否", "MEMBER", "MEMBER；负责人请改项目", "member_role"),
    ],
    "resource_types": [
        ("code", "是", "CURSOR", "类型编码，唯一", "code"),
        ("name", "是", "Cursor", "类型名称", "name"),
        ("category", "否", "AI", "资源类别", "category"),
        ("description", "否", "", "说明", "description"),
    ],
    "resources": [
        ("resource_name", "是", "Cursor 主账号", "资源名称", "resource_name"),
        ("resource_type_code", "是", "CURSOR", "已存在资源类型", "resource_type_code"),
        ("provider", "否", "Cursor", "供应商", "provider"),
        ("account_identifier", "否", "team@example.com", "账号标识，不是密码", "account_identifier"),
        ("plan_name", "否", "Pro", "套餐", "plan_name"),
        ("status", "否", "ACTIVE", "ACTIVE / IDLE / EXPIRED / DISABLED", "status"),
        ("purchased_at", "否", "2026-01-01", "YYYY-MM-DD", "purchased_at"),
        ("activated_at", "否", "2026-01-02", "YYYY-MM-DD", "activated_at"),
        ("expires_at", "否", "2026-12-31", "YYYY-MM-DD", "expires_at"),
        ("remark", "否", "", "备注", "remark"),
    ],
    "resource_allocations": [
        ("resource_name", "是", "Cursor 主账号", "已存在资源名称", "resource_name"),
        ("project_code", "否", "P-001", "项目与成员至少一项", "project_code"),
        ("employee_no", "否", "E1001", "项目+成员时必须是在项成员", "employee_no"),
        ("remark", "否", "", "备注", "remark"),
    ],
    "expenses": [
        ("expense_type_code", "是", "SAAS", "已存在费用类型", "expense_type_code"),
        ("resource_name", "否", "Cursor 主账号", "可选绑定资源", "resource_name"),
        ("billing_mode", "是", "MANUAL", "RECURRING / ONE_TIME / USAGE / MANUAL", "billing_mode"),
        ("expense_date", "是", "2026-09-01", "YYYY-MM-DD", "expense_date"),
        ("amount", "是", "100.00", "数字，独立币种列", "amount"),
        ("currency", "是", "CNY", "CNY / USD", "currency"),
        ("description", "否", "九月订阅", "说明", "description"),
        ("project_code", "否", "P-001", "项目与成员至少一项", "project_code"),
        ("employee_no", "否", "E1001", "项目+成员时必须是在项成员", "employee_no"),
        ("allocation_method", "否", "DIRECT", "DIRECT / EQUAL / RATIO / MANUAL", "allocation_method"),
        ("allocated_amount", "否", "100.00", "空则等于费用金额", "allocated_amount"),
        ("billing_period_start", "否", "2026-09-01", "YYYY-MM-DD", "billing_period_start"),
        ("billing_period_end", "否", "2026-09-30", "YYYY-MM-DD", "billing_period_end"),
        ("external_reference", "否", "INV-001", "外部单号", "external_reference"),
    ],
}

HEADER_ALIASES: dict[str, dict[str, str]] = {
    "employees": {
        "工号": "employee_no",
        "姓名": "name",
        "部门": "department",
        "职位": "position",
        "邮箱": "email",
        "电话": "phone",
        "在职状态": "employment_status",
        "备注": "remark",
    },
    "projects": {
        "项目编码": "project_code",
        "项目名称": "name",
        "负责人工号": "manager_employee_no",
        "状态": "status",
        "说明": "description",
        "开始日期": "start_date",
        "结束日期": "end_date",
    },
    "project_members": {
        "项目编码": "project_code",
        "工号": "employee_no",
        "成员角色": "member_role",
    },
    "resource_types": {
        "编码": "code",
        "名称": "name",
        "类别": "category",
        "说明": "description",
    },
    "resources": {
        "资源名称": "resource_name",
        "资源类型编码": "resource_type_code",
        "供应商": "provider",
        "账号标识": "account_identifier",
        "套餐": "plan_name",
        "状态": "status",
        "采购日期": "purchased_at",
        "启用日期": "activated_at",
        "到期日期": "expires_at",
        "备注": "remark",
    },
    "resource_allocations": {
        "资源名称": "resource_name",
        "项目编码": "project_code",
        "工号": "employee_no",
        "备注": "remark",
    },
    "expenses": {
        "费用类型编码": "expense_type_code",
        "资源名称": "resource_name",
        "计费方式": "billing_mode",
        "费用日期": "expense_date",
        "金额": "amount",
        "币种": "currency",
        "说明": "description",
        "项目编码": "project_code",
        "工号": "employee_no",
        "分摊方式": "allocation_method",
        "分摊金额": "allocated_amount",
        "账期开始": "billing_period_start",
        "账期结束": "billing_period_end",
        "外部单号": "external_reference",
    },
}


def ensure_import_type(import_type: str) -> str:
    if import_type not in TEMPLATE_COLUMNS:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="不支持的导入类型")
    return import_type


def validate_upload(filename: str | None, size: int | None) -> None:
    name = (filename or "").strip()
    if not name.lower().endswith(".xlsx") or name.lower().endswith(".xlsm"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持 .xlsx 文件")
    if ".." in name.replace("\\", "/").split("/")[-1] and name.startswith(("..", "/", "\\")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="非法文件名")
    if "/" in name or "\\" in name:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="非法文件名")
    if size is not None and size > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件超过 10 MB")


def _header_key(import_type: str, raw: Any) -> str | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    lowered = text.lower()
    if lowered in FORBIDDEN_HEADERS or any(token in lowered for token in FORBIDDEN_HEADERS):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"不允许导入凭据字段: {text}",
        )
    aliases = HEADER_ALIASES.get(import_type, {})
    if text in aliases:
        return aliases[text]
    if lowered in aliases:
        return aliases[lowered]
    known = {field for _, _, _, _, field in TEMPLATE_COLUMNS[import_type]}
    if lowered in known:
        return lowered
    return lowered


def _is_formula(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).lstrip()
    if not text:
        return False
    if text[0] in {"=", "+", "@"}:
        return True
    return False


def _cell_to_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(str(value)), "f")
    text = str(value).strip()
    return text or None


def parse_workbook(import_type: str, content: bytes) -> list[tuple[int, dict[str, str]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件超过 10 MB")
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="无法读取 Excel 文件") from exc
    try:
        if "数据" in workbook.sheetnames:
            sheet = workbook["数据"]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="模板缺少表头") from exc
        headers: list[str | None] = []
        seen: set[str] = set()
        for cell in header_row:
            key = _header_key(import_type, cell)
            if key:
                if key in seen:
                    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"重复列: {key}")
                seen.add(key)
            headers.append(key)
        known = {field for _, _, _, _, field in TEMPLATE_COLUMNS[import_type]}
        if not any(key in known for key in headers if key):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="未识别到有效列")
        parsed: list[tuple[int, dict[str, str]]] = []
        excel_row = 1
        for raw in rows_iter:
            excel_row += 1
            if excel_row - 1 > MAX_IMPORT_ROWS:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="单批次最多 10000 行",
                )
            values: dict[str, str] = {}
            empty = True
            for index, cell in enumerate(raw):
                if index >= len(headers) or not headers[index]:
                    continue
                if _is_formula(cell):
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"第 {excel_row} 行存在公式，已拒绝",
                    )
                text = _cell_to_str(cell)
                if text is not None:
                    empty = False
                    values[headers[index] or ""] = text
            if empty:
                continue
            parsed.append((excel_row, values))
        return parsed
    finally:
        workbook.close()


def build_template(import_type: str) -> bytes:
    ensure_import_type(import_type)
    columns = TEMPLATE_COLUMNS[import_type]
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "数据"
    header_font = Font(bold=True, color="161616", name="IBM Plex Sans")
    fill = PatternFill("solid", fgColor="F4F4F4")
    for index, (header, _required, example, _note, _field) in enumerate(columns, start=1):
        cell = data_sheet.cell(row=1, column=index, value=header)
        cell.font = header_font
        cell.fill = fill
        data_sheet.cell(row=2, column=index, value=example)
        data_sheet.column_dimensions[get_column_letter(index)].width = 22
    help_sheet = workbook.create_sheet("说明")
    help_headers = ["字段名", "是否必填", "示例", "说明"]
    for index, title in enumerate(help_headers, start=1):
        cell = help_sheet.cell(row=1, column=index, value=title)
        cell.font = header_font
        cell.fill = fill
    for row_index, (header, required, example, note, _field) in enumerate(columns, start=2):
        help_sheet.cell(row=row_index, column=1, value=header)
        help_sheet.cell(row=row_index, column=2, value=required)
        help_sheet.cell(row=row_index, column=3, value=example)
        help_sheet.cell(row=row_index, column=4, value=note)
        help_sheet.cell(row=row_index, column=4).alignment = Alignment(wrap_text=True)
    for index in range(1, 5):
        help_sheet.column_dimensions[get_column_letter(index)].width = 28
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def new_export_workbook(title: str) -> tuple[Workbook, Worksheet]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "export"
    return workbook, sheet


def write_header(sheet: Worksheet, headers: list[str]) -> None:
    font = Font(bold=True, color="161616")
    fill = PatternFill("solid", fgColor="F4F4F4")
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = font
        cell.fill = fill
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(32, len(header) + 8))


def write_cell(sheet: Worksheet, row: int, column: int, value: Any, *, money: bool = False, as_date: bool = False) -> None:
    cell = sheet.cell(row=row, column=column)
    if value is None:
        cell.value = None
        return
    if money:
        cell.value = float(Decimal(str(value)))
        cell.number_format = "0.00##"
        return
    if as_date or isinstance(value, date):
        if isinstance(value, datetime):
            cell.value = value.date()
        elif isinstance(value, date):
            cell.value = value
        else:
            cell.value = str(value)
        cell.number_format = "YYYY-MM-DD"
        return
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        cell.value = f"'{value}"
        return
    cell.value = value


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
