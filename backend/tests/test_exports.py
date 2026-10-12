from io import BytesIO

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from tests.helpers import (
    admin_headers,
    auth_headers_for,
    create_employee,
    create_expense,
    create_project,
    create_resource,
    create_user_account,
    unique,
)


def _xlsx_rows(content: bytes) -> list[list[object]]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_admin_global_export(client: TestClient) -> None:
    headers = admin_headers(client)
    create_project(client, headers, project_code=unique("EXA"))
    response = client.get("/api/exports/projects", headers=headers)
    assert response.status_code == 200
    rows = _xlsx_rows(response.content)
    assert rows[0][1] == "project_code"
    assert any(row[1] and str(row[1]).startswith("EXA") for row in rows[1:])


def test_boss_global_export(client: TestClient) -> None:
    admin = admin_headers(client)
    employee = create_employee(client, admin)
    _, username, password = create_user_account(client, admin, role_code="BOSS", employee_id=employee["id"])
    headers = auth_headers_for(client, username, password)
    response = client.get("/api/exports/projects", headers=headers)
    assert response.status_code == 200


def test_pm_only_managed_projects(client: TestClient) -> None:
    admin = admin_headers(client)
    manager = create_employee(client, admin)
    other = create_employee(client, admin)
    mine = create_project(client, admin, manager_employee_id=manager["id"], project_code=unique("MINE"))
    create_project(client, admin, manager_employee_id=other["id"], project_code=unique("OTHR"))
    _, username, password = create_user_account(
        client, admin, role_code="PROJECT_MANAGER", employee_id=manager["id"]
    )
    headers = auth_headers_for(client, username, password)
    response = client.get("/api/exports/projects", headers=headers)
    assert response.status_code == 200
    codes = [str(row[1]) for row in _xlsx_rows(response.content)[1:] if row[1]]
    assert mine["project_code"] in codes
    assert all(not code.startswith("OTHR") for code in codes)


def test_member_only_self_expenses(client: TestClient) -> None:
    admin = admin_headers(client)
    member_emp = create_employee(client, admin)
    other_emp = create_employee(client, admin)
    project = create_project(client, admin)
    assert client.post(
        f"/api/projects/{project['id']}/members",
        headers=admin,
        json={"employee_id": member_emp["id"], "member_role": "MEMBER"},
    ).status_code == 200
    assert client.post(
        f"/api/projects/{project['id']}/members",
        headers=admin,
        json={"employee_id": other_emp["id"], "member_role": "MEMBER"},
    ).status_code == 200
    create_expense(
        client,
        admin,
        distributions=[{"project_id": project["id"], "employee_id": member_emp["id"], "allocation_method": "DIRECT", "allocated_amount": "40.00"}],
        description="member-export",
        amount="40.00",
        currency="CNY",
    )
    create_expense(
        client,
        admin,
        distributions=[{"project_id": project["id"], "employee_id": other_emp["id"], "allocation_method": "DIRECT", "allocated_amount": "90.00"}],
        description="other-export",
        amount="90.00",
        currency="USD",
    )
    _, username, password = create_user_account(client, admin, role_code="MEMBER", employee_id=member_emp["id"])
    headers = auth_headers_for(client, username, password)
    response = client.get("/api/exports/expenses", headers=headers)
    assert response.status_code == 200
    rows = _xlsx_rows(response.content)[1:]
    descriptions = [row[3] for row in rows]
    assert "member-export" in descriptions
    assert "other-export" not in descriptions
    currencies = {row[5] for row in rows}
    assert "USD" not in currencies


def test_export_keeps_currency_column(client: TestClient) -> None:
    admin = admin_headers(client)
    project = create_project(client, admin)
    create_expense(
        client,
        admin,
        distributions=[{"project_id": project["id"], "allocation_method": "DIRECT", "allocated_amount": "10.00"}],
        amount="10.00",
        currency="CNY",
        description="cny-only",
    )
    create_expense(
        client,
        admin,
        distributions=[{"project_id": project["id"], "allocation_method": "DIRECT", "allocated_amount": "11.00"}],
        amount="11.00",
        currency="USD",
        description="usd-only",
    )
    response = client.get("/api/exports/expenses", headers=admin)
    rows = _xlsx_rows(response.content)
    assert rows[0][4] == "amount"
    assert rows[0][5] == "currency"
    pairs = {(row[3], row[5]) for row in rows[1:]}
    assert ("cny-only", "CNY") in pairs
    assert ("usd-only", "USD") in pairs
    amounts = [row[4] for row in rows[1:] if row[3] in {"cny-only", "usd-only"}]
    assert not any(isinstance(value, str) and "CNY" in value for value in amounts)


def test_voided_excluded_from_expense_export(client: TestClient) -> None:
    admin = admin_headers(client)
    project = create_project(client, admin)
    expense = create_expense(
        client,
        admin,
        distributions=[{"project_id": project["id"], "allocation_method": "DIRECT", "allocated_amount": "12.00"}],
        description="to-void",
        amount="12.00",
    )
    voided = client.post(f"/api/expenses/{expense['id']}/void", headers=admin)
    assert voided.status_code == 200
    response = client.get("/api/exports/expenses", headers=admin)
    descriptions = [row[3] for row in _xlsx_rows(response.content)[1:]]
    assert "to-void" not in descriptions


def test_member_cannot_export_audit(client: TestClient) -> None:
    admin = admin_headers(client)
    employee = create_employee(client, admin)
    _, username, password = create_user_account(client, admin, role_code="MEMBER", employee_id=employee["id"])
    headers = auth_headers_for(client, username, password)
    response = client.get("/api/exports/audit_logs", headers=headers)
    assert response.status_code == 403
